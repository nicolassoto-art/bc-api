"""Alta de unidades por Excel: celda vacía de descuento/bono = "sin dato" (None), no 0.

Un 0 guardado tapa el bono de la ficha del proyecto en el editor y en el cotizador;
None hace que caigan a la ficha. Solo un 0 explícito en el Excel es un cero real.
"""
import io
import uuid

from openpyxl import Workbook

from app.models import Proyecto, Unidad

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _excel(filas):
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("INSTRUCCIONES")
    u = wb.create_sheet("UNIDAD")
    u.append(["REQ", "REQ", "REQ", "OPC", "OPC"])
    u.append(["Unidad Número", "Modelo", "ValorUF", "Descuento", "Bonopie"])
    for fila in filas:
        u.append(list(fila))
    for hoja in ("ESTACIONAMIENTOS", "BODEGAS"):
        s = wb.create_sheet(hoja)
        s.append(["OPC", "OPC"])
        s.append(["Número", "PrecioUF"])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def _subir(client, h, pid, filas):
    r = client.post(f"/proyectos/{pid}/unidades/excel/upload", headers=h,
                    files={"file": ("x.xlsx", _excel(filas), XLSX)})
    assert r.status_code == 200, r.text


def _unidades(db, pid):
    db.expire_all()
    return {u.numero: u for u in db.query(Unidad).filter(Unidad.proyecto_id == pid)}


def test_alta_celda_vacia_queda_sin_dato_y_cero_explicito_se_respeta(client, admin, db):
    _, h = admin
    pid = "test-alta-" + uuid.uuid4().hex[:8]
    db.add(Proyecto(id=pid, nombre=pid, extra={}))
    db.commit()

    _subir(client, h, pid, [
        ("101", "2D1B", 2500, None, None),  # la fuente no trae descuento/bono
        ("102", "2D1B", 2500, 0, 0),        # cero real informado por la fuente
        ("103", "2D1B", 2500, 5, 10),
    ])

    u = _unidades(db, pid)
    assert u["101"].descuento_pct is None and u["101"].bono_pie_pct is None
    assert u["101"].precio_final_uf == 2500
    assert u["102"].descuento_pct == 0 and u["102"].bono_pie_pct == 0
    assert u["103"].descuento_pct == 5 and u["103"].bono_pie_pct == 10
    assert u["103"].precio_final_uf == 2375


def test_unidad_existente_conserva_bono_manual_con_celda_vacia(client, admin, db):
    _, h = admin
    pid = "test-alta-" + uuid.uuid4().hex[:8]
    db.add(Proyecto(id=pid, nombre=pid, extra={}))
    db.commit()
    db.add(Unidad(id=f"{pid}-201", proyecto_id=pid, numero="201", tipo="Depto",
                  modelo="2D1B", disponible=True, bono_pie_pct=10, descuento_pct=3))
    db.commit()

    _subir(client, h, pid, [("201", "2D1B", 2600, None, None)])

    u = _unidades(db, pid)["201"]
    assert u.bono_pie_pct == 10 and u.descuento_pct == 3
