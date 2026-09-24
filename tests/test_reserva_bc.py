"""Marca "reservada por BigCapital" (reserva_bc) — 2026-09-24.

Nuestro stock y el de la inmobiliaria son dos. Al reservar en la intranet la unidad
sale del nuestro, pero sigue libre en PlanOk hasta que la inmobiliaria la registra;
la lectura horaria ("aparece en el Excel = disponible") la volvía a abrir. Con la
marca, ningún camino puede dejarla disponible hasta que la reserva se anule.
"""
import io
import uuid

from openpyxl import Workbook

from app.models import Proyecto, Unidad

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _pid():
    return "test-rbc-" + uuid.uuid4().hex[:8]


def _excel(numeros):
    """Excel JB mínimo con las unidades dadas (todas 'aparecen' = disponibles)."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("INSTRUCCIONES")
    u = wb.create_sheet("UNIDAD")
    u.append(["REQ", "REQ", "REQ"])
    u.append(["Unidad Número", "Modelo", "ValorUF"])
    for n in numeros:
        u.append([n, "2D1B", 2500])
    # El formato JB exige las 4 hojas; sin estas cae al parser legado.
    for hoja in ("ESTACIONAMIENTOS", "BODEGAS"):
        s = wb.create_sheet(hoja)
        s.append(["OPC", "OPC"])
        s.append(["Número", "PrecioUF"])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def _subir(client, h, pid, numeros):
    r = client.post(f"/proyectos/{pid}/unidades/excel/upload", headers=h,
                    files={"file": ("x.xlsx", _excel(numeros), XLSX)})
    assert r.status_code == 200, r.text
    return r.json()


def _proyecto_con(db, pid, unidades):
    """unidades: {numero: disponible}"""
    db.add(Proyecto(id=pid, nombre=pid, extra={}))
    db.commit()
    for n, disp in unidades.items():
        db.add(Unidad(id=f"{pid}-{n}", proyecto_id=pid, numero=n, tipo="Depto",
                      modelo="2D1B", disponible=disp))
    db.commit()


def _unidad(db, pid, numero):
    db.expire_all()
    return db.query(Unidad).filter(Unidad.proyecto_id == pid, Unidad.numero == numero).one()


def _marcar(client, h, pid, uid, **body):
    return client.put(f"/proyectos/{pid}/unidades/{uid}/reserva-bc", headers=h, json=body)


def test_marcar_saca_del_stock_y_avisa_como_estaba(client, admin, db):
    _, h = admin
    pid = _pid()
    _proyecto_con(db, pid, {"101": True})
    r = _marcar(client, h, pid, f"{pid}-101", reserva_bc="RES-1")
    assert r.status_code == 200, r.text
    assert r.json()["estaba_disponible"] is True
    assert r.json()["unidad"]["reserva_bc"] == "RES-1"
    u = _unidad(db, pid, "101")
    assert u.disponible is False and u.reserva_bc == "RES-1" and u.reserva_bc_at is not None


def test_marcar_unidad_que_ya_no_estaba_disponible(client, admin, db):
    _, h = admin
    pid = _pid()
    _proyecto_con(db, pid, {"101": False})
    r = _marcar(client, h, pid, f"{pid}-101", reserva_bc="RES-1")
    assert r.status_code == 200 and r.json()["estaba_disponible"] is False
    # Liberar sin reabrir (esta reserva no la sacó): sigue no disponible, sin marca.
    r = _marcar(client, h, pid, f"{pid}-101", reserva_bc=None, esperado="RES-1", reabrir=False)
    assert r.status_code == 200
    u = _unidad(db, pid, "101")
    assert u.reserva_bc is None and u.disponible is False


def test_excel_no_reabre_una_unidad_marcada(client, admin, db):
    """El escenario real: reservamos y la lectura horaria la trae libre desde PlanOk."""
    _, h = admin
    pid = _pid()
    _proyecto_con(db, pid, {"101": True})
    assert _marcar(client, h, pid, f"{pid}-101", reserva_bc="RES-1").status_code == 200
    _subir(client, h, pid, ["101", "102"])
    u = _unidad(db, pid, "101")
    assert u.disponible is False and u.reserva_bc == "RES-1"
    assert _unidad(db, pid, "102").disponible is True  # la nueva, sin marca, normal


def test_excel_no_anota_un_cambio_falso_en_el_timeline(client, admin, db):
    """Segunda subida idéntica con la unidad marcada: debe quedar 'Sin cambios'."""
    _, h = admin
    pid = _pid()
    db.add(Proyecto(id=pid, nombre=pid, extra={}))
    db.commit()
    _subir(client, h, pid, ["101"])  # la crea tal como viene en el Excel
    uid = _unidad(db, pid, "101").id
    assert _marcar(client, h, pid, uid, reserva_bc="RES-1").status_code == 200
    _subir(client, h, pid, ["101"])
    db.expire_all()
    proy = db.query(Proyecto).filter(Proyecto.id == pid).one()
    ultimo = (proy.extra or {}).get("timeline", [])[0]
    assert "Sin cambios" in ultimo["detalles"], ultimo["detalles"]


def test_excel_sin_la_unidad_marcada_la_deja_no_disponible(client, admin, db):
    _, h = admin
    pid = _pid()
    _proyecto_con(db, pid, {"101": True, "102": True})
    assert _marcar(client, h, pid, f"{pid}-101", reserva_bc="RES-1").status_code == 200
    _subir(client, h, pid, ["102"])
    u = _unidad(db, pid, "101")
    assert u.disponible is False and u.reserva_bc == "RES-1"


def test_unidad_sin_marca_se_comporta_como_siempre(client, admin, db):
    _, h = admin
    pid = _pid()
    _proyecto_con(db, pid, {"101": False, "102": True})
    _subir(client, h, pid, ["101"])
    assert _unidad(db, pid, "101").disponible is True   # aparece en Excel = disponible
    assert _unidad(db, pid, "102").disponible is False  # no viene = baja


def test_quitar_la_marca_y_reimportar_la_vuelve_a_abrir(client, admin, db):
    _, h = admin
    pid = _pid()
    _proyecto_con(db, pid, {"101": True})
    assert _marcar(client, h, pid, f"{pid}-101", reserva_bc="RES-1").status_code == 200
    r = _marcar(client, h, pid, f"{pid}-101", reserva_bc=None, esperado="RES-1", reabrir=True)
    assert r.status_code == 200
    u = _unidad(db, pid, "101")
    assert u.reserva_bc is None and u.disponible is True
    _subir(client, h, pid, ["101"])
    assert _unidad(db, pid, "101").disponible is True


def test_una_reserva_no_libera_la_marca_de_otra(client, admin, db):
    _, h = admin
    pid = _pid()
    _proyecto_con(db, pid, {"101": True})
    assert _marcar(client, h, pid, f"{pid}-101", reserva_bc="RES-1").status_code == 200
    r = _marcar(client, h, pid, f"{pid}-101", reserva_bc=None, esperado="RES-2", reabrir=True)
    assert r.status_code == 409
    u = _unidad(db, pid, "101")
    assert u.reserva_bc == "RES-1" and u.disponible is False


def test_no_se_marca_una_unidad_marcada_por_otra_reserva(client, admin, db):
    _, h = admin
    pid = _pid()
    _proyecto_con(db, pid, {"101": True})
    assert _marcar(client, h, pid, f"{pid}-101", reserva_bc="RES-1").status_code == 200
    assert _marcar(client, h, pid, f"{pid}-101", reserva_bc="RES-2").status_code == 409
    assert _unidad(db, pid, "101").reserva_bc == "RES-1"
    # Repetir la misma marca es idempotente
    assert _marcar(client, h, pid, f"{pid}-101", reserva_bc="RES-1").status_code == 200


def test_el_put_del_editor_no_borra_la_marca_ni_la_reabre(client, admin, db):
    """El editor reenvía la unidad completa desde una copia que puede ser vieja."""
    _, h = admin
    pid = _pid()
    _proyecto_con(db, pid, {"101": True})
    assert _marcar(client, h, pid, f"{pid}-101", reserva_bc="RES-1").status_code == 200
    copia_vieja = {"numero": "101", "modelo": "2D1B", "tipo": "Depto",
                   "disponible": True, "reserva_bc": None}
    r = client.put(f"/proyectos/{pid}/unidades/{pid}-101", headers=h, json=copia_vieja)
    assert r.status_code == 200, r.text
    u = _unidad(db, pid, "101")
    assert u.reserva_bc == "RES-1" and u.disponible is False
    assert r.json()["reserva_bc"] == "RES-1"


def test_la_salida_de_unidades_muestra_la_marca(client, admin, db):
    _, h = admin
    pid = _pid()
    _proyecto_con(db, pid, {"101": True})
    assert _marcar(client, h, pid, f"{pid}-101", reserva_bc="RES-9").status_code == 200
    r = client.get(f"/proyectos/{pid}/unidades", headers=h)
    assert r.status_code == 200
    fila = next(x for x in r.json() if x["numero"] == "101")
    assert fila["reserva_bc"] == "RES-9" and fila["disponible"] is False


def test_reimportacion_con_overwrite_conserva_la_marca(client, admin, db):
    """El importador por lotes borra en bloque y recrea: la marca debe sobrevivir."""
    _, h = admin
    jb_id = "rbc" + uuid.uuid4().hex[:6]
    proyecto_jb = {"id": jb_id, "name": "Reimport Marca",
                   "units": [{"number": "101", "available": True},
                             {"number": "102", "available": True}]}
    r = client.post("/importador/batch", headers=h, json={"projects": [proyecto_jb]})
    assert r.status_code == 200, r.text
    pid = r.json()["details"][0]["project_id"]
    uid = _unidad(db, pid, "101").id
    assert _marcar(client, h, pid, uid, reserva_bc="RES-7").status_code == 200
    r = client.post("/importador/batch", headers=h,
                    json={"projects": [proyecto_jb], "overwrite": True})
    assert r.status_code == 200, r.text
    u = _unidad(db, pid, "101")
    assert u.reserva_bc == "RES-7" and u.disponible is False
    assert _unidad(db, pid, "102").disponible is True


def test_la_red_final_cubre_cualquier_escritura(db):
    """Aunque un camino nuevo olvide la regla, el guardado la hace cumplir."""
    pid = _pid()
    _proyecto_con(db, pid, {"101": True})
    u = _unidad(db, pid, "101")
    u.reserva_bc = "RES-1"
    db.commit()
    u.disponible = True  # un importador cualquiera intenta abrirla
    db.commit()
    assert _unidad(db, pid, "101").disponible is False
