"""Unidad NUEVA por Excel: descuento/bono con celda vacía quedan None ("sin dato"), no 0.

Regla del negocio (2026-08-01): null = la unidad no tiene dato propio y hereda el % de
la ficha del proyecto; 0 = cero real de ESA unidad, que tapa la ficha. bc-api inventaba
el 0 en cada alta (default del modelo + default explícito en subir_excel) y el editor lo
tomaba como dato real: al 2026-09-30 había 239 unidades de Euro cotizando bono pie 0%
con la ficha en 5-15%, porque Mobysuite no entrega bono y el scraper manda la celda vacía.
"""
import io
import uuid

from openpyxl import Workbook

from app.models import Proyecto, Unidad

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _excel(filas):
    """Excel JB mínimo con columnas Descuento y Bonopie en la hoja UNIDAD."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("INSTRUCCIONES")
    u = wb.create_sheet("UNIDAD")
    u.append(["REQ", "REQ", "REQ", "OPC", "OPC"])
    u.append(["Unidad Número", "Modelo", "ValorUF", "Descuento", "Bonopie"])
    for f in filas:
        u.append(f)
    for hoja in ("ESTACIONAMIENTOS", "BODEGAS"):
        s = wb.create_sheet(hoja)
        s.append(["OPC", "OPC"])
        s.append(["Número", "PrecioUF"])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def _proy(db, base):
    """Proyecto NUEVO por corrida: con un id fijo, una corrida anterior contra la misma
    BD deja la unidad creada y el alta se vuelve actualización (lo que este test no mide)."""
    pid = f"{base}-{uuid.uuid4().hex[:8]}"
    db.add(Proyecto(id=pid, nombre=pid, extra={}))
    db.commit()
    return pid


def _subir(client, h, pid, filas):
    r = client.post(f"/proyectos/{pid}/unidades/excel/upload", headers=h,
                    files={"file": ("x.xlsx", _excel(filas), XLSX)})
    assert r.status_code == 200, r.text
    return r.json()


def _unidad(db, pid, numero):
    db.expire_all()
    return db.query(Unidad).filter(Unidad.proyecto_id == pid, Unidad.numero == numero).one()


def test_alta_con_celdas_vacias_queda_sin_dato(client, admin, db):
    _, h = admin
    pid = _proy(db, "test-alta-sin-dato")
    _subir(client, h, pid, [["101", "2D1B", 2500, None, None]])
    u = _unidad(db, pid, "101")
    assert u.descuento_pct is None
    assert u.bono_pie_pct is None
    assert u.precio_final_uf == 2500  # sin descuento propio: final = lista


def test_alta_con_cero_explicito_se_respeta(client, admin, db):
    """Un 0 que SÍ viene en el Excel es dato real de la fuente: se guarda 0."""
    _, h = admin
    pid = _proy(db, "test-alta-cero-real")
    _subir(client, h, pid, [["101", "2D1B", 2500, 0, 0]])
    u = _unidad(db, pid, "101")
    assert u.descuento_pct == 0
    assert u.bono_pie_pct == 0


def test_alta_con_valor_real_se_respeta(client, admin, db):
    _, h = admin
    pid = _proy(db, "test-alta-valor-real")
    _subir(client, h, pid, [["101", "2D1B", 2500, 7, 10]])
    u = _unidad(db, pid, "101")
    assert u.descuento_pct == 7
    assert u.bono_pie_pct == 10
    assert u.precio_final_uf == 2325  # 2500 × (1 − 7%)


def test_resubida_con_celdas_vacias_no_pisa_el_dato_existente(client, admin, db):
    """Regresión del camino de actualización: vacío sigue preservando lo cargado."""
    _, h = admin
    pid = _proy(db, "test-resubida-preserva")
    _subir(client, h, pid, [["101", "2D1B", 2500, 7, 10]])
    _subir(client, h, pid, [["101", "2D1B", 2500, None, None]])
    u = _unidad(db, pid, "101")
    assert u.descuento_pct == 7
    assert u.bono_pie_pct == 10
