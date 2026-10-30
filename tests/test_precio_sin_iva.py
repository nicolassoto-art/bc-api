"""Precio normal (con IVA) y precio sin IVA en MNK — 2026-10-29.

En MNK el precio del negocio es el SIN IVA (casilla «Descontar IVA» de PlanOk). Con el
interruptor del scraper activo, el Excel trae `ValorUF` = sin IVA y una columna nueva
al final: «PrecioConIvaUF» (precio normal, solo para mostrar). No existe un estado
«precio en revisión»: toda unidad publicada trae su par normal + sin IVA.

Regla dura: si esa columna NO viene, todo queda idéntico a antes (otras
inmobiliarias, MNK sin activar, un Excel viejo o una restauración por curl).
"""
import copy
import io
import uuid

from openpyxl import Workbook

from app.models import Proyecto, Unidad
from app.routes.proyectos import _unidad_dict
from app.routes.unidades import _build_jb_extras

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NOMBRE = "Prueba Interna"


def _excel(filas, cols=("Unidad Número", "Modelo", "ValorUF"), estac=None, bodegas=None,
           cols_sec=("Número", "PrecioUF")):
    """Excel JB mínimo. `filas`/`estac`/`bodegas` son listas alineadas con sus columnas."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("INSTRUCCIONES")
    u = wb.create_sheet("UNIDAD")
    u.append(["REQ"] * len(cols))
    u.append(list(cols))
    for f in filas:
        u.append(list(f))
    for hoja, datos in (("ESTACIONAMIENTOS", estac), ("BODEGAS", bodegas)):
        s = wb.create_sheet(hoja)
        s.append(["OPC"] * len(cols_sec))
        s.append(list(cols_sec))
        for f in datos or []:
            s.append(list(f))
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def _proy(db, base="test-siniva", extra=None):
    pid = f"{base}-{uuid.uuid4().hex[:8]}"
    db.add(Proyecto(id=pid, nombre=f"{NOMBRE} {pid}", inmobiliaria="MNK", extra=extra or {}))
    db.commit()
    return pid


def _subir(client, h, pid, xlsx, origen=None):
    url = f"/proyectos/{pid}/unidades/excel/upload" + (f"?origen={origen}" if origen else "")
    r = client.post(url, headers=h, files={"file": ("x.xlsx", xlsx, XLSX)})
    assert r.status_code == 200, r.text
    return r.json()


def _unidad(db, pid, numero):
    db.expire_all()
    return db.query(Unidad).filter(Unidad.proyecto_id == pid, Unidad.numero == numero).one()


def _proyecto(db, pid):
    db.expire_all()
    return db.get(Proyecto, pid)


COLS_SIN_IVA = ("Unidad Número", "Modelo", "ValorUF", "PrecioConIvaUF")
SEC_SIN_IVA = ("Número", "PrecioUF", "PrecioConIvaUF")


# ── Sin las columnas nuevas: idéntico a antes ──────────────────────────────────

def test_sin_columnas_la_unidad_queda_igual_que_siempre(client, admin, db):
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 3168]]))
    u = _unidad(db, pid, "16B")
    assert u.precio_lista_uf == 3168 and u.precio_final_uf == 3168
    assert u.precio_con_iva_uf is None
    # La salida pública de la unidad NO suma claves nuevas
    d = _unidad_dict(u)
    assert "precio_con_iva_uf" not in d


def test_build_jb_extras_sin_columnas_identico():
    """Salida exacta de siempre (claves y valores) cuando la hoja no trae las columnas."""
    estac = [{"Número": "E-1A", "PrecioUF": 215, "Nivel": "-1", "Tipo": "Simple"},
             {"Número": "E-2A", "PrecioUF": 200, "Pack Número": "P1"}]
    bod = [{"Número": "B-1A", "PrecioUF": 80, "Superficie": 3.5, "Pack Número": "P1"}]
    out = _build_jb_extras(copy.deepcopy(estac), copy.deepcopy(bod))
    assert out["estacionamientos"][0] == {
        "id": "jb-est-0", "numero": "E-1A", "precio_uf": 215.0, "nivel": "-1",
        "tipo": "Simple", "tenencia": "", "disponible": True}
    assert out["bodegas"][0] == {
        "id": "jb-bod-0", "numero": "B-1A", "precio_uf": 80.0, "superficie": 3.5, "tipo": "",
        "tenencia": "", "disponible": False, "pack_numero": "P1"}
    assert out["packs"] == [{"id": "jb-pack-0", "numero": "P1", "estacionamientos": ["E-2A"],
                             "bodegas": ["B-1A"], "precio_uf": 280.0, "disponible": True}]
    assert out["_estacionamientos_dom"] == estac and out["_bodegas_dom"] == bod


def test_sin_columnas_cambio_de_estacionamiento_no_genera_evento_nuevo(client, admin, db):
    """Hoy el timeline no mira estacionamientos/bodegas: sin precio normal, sigue igual."""
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 3168]], estac=[["E-1A", 215]]), origen="jb_importer")
    antes = len(_proyecto(db, pid).extra.get("timeline") or [])
    _subir(client, h, pid, _excel([["16B", "B1", 3168]], estac=[["E-1A", 180]]), origen="jb_importer")
    p = _proyecto(db, pid)
    assert len(p.extra.get("timeline") or []) == antes  # corrida automática «sin cambios»
    assert p.extra["estacionamientos"][0]["precio_uf"] == 180
    assert "precio_con_iva_uf" not in p.extra["estacionamientos"][0]


def test_listado_sin_dato_trae_desde_con_iva_nulo(client, admin, db):
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 3168], ["17B", "B1", 3000]]))
    r = client.get("/proyectos", headers=h)
    assert r.status_code == 200
    fila = next(x for x in r.json() if x["id"] == pid)
    assert fila["precio_desde_uf"] == 3000
    assert fila["precio_desde_con_iva_uf"] is None


# ── Con las columnas ──────────────────────────────────────────────────────────

def test_con_columnas_guarda_precio_normal(client, admin, db):
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel(
        [["16B", "B1", 2704.56, 3168], ["17B", "B1", 2561, 3000]],
        cols=COLS_SIN_IVA,
        estac=[["E-1A", 173.11, 215]], cols_sec=SEC_SIN_IVA))
    u = _unidad(db, pid, "16B")
    assert u.precio_lista_uf == 2704.56 and u.precio_con_iva_uf == 3168
    d = _unidad_dict(u)
    assert d["precio_con_iva_uf"] == 3168 and "precio_en_revision" not in d
    p = _proyecto(db, pid)
    e = p.extra["estacionamientos"][0]
    assert e["precio_uf"] == 173.11 and e["precio_con_iva_uf"] == 215
    # GET /unidades expone los campos
    r = client.get(f"/proyectos/{pid}/unidades", headers=h)
    fila = next(x for x in r.json() if x["numero"] == "16B")
    assert fila["precio_con_iva_uf"] == 3168 and "precio_en_revision" not in fila


def test_packs_suman_precio_normal_solo_si_todos_lo_traen():
    estac = [{"Número": "E-2A", "PrecioUF": 170, "PrecioConIvaUF": 200, "Pack Número": "P1"}]
    bod = [{"Número": "B-1A", "PrecioUF": 70, "PrecioConIvaUF": 80.5, "Pack Número": "P1"}]
    out = _build_jb_extras(estac, bod)
    assert out["packs"][0]["precio_con_iva_uf"] == 280.5
    # Bodega sin el dato (hoja sin la columna) → el pack no inventa una suma parcial
    bod2 = [{"Número": "B-1A", "PrecioUF": 70, "Pack Número": "P1"}]
    out2 = _build_jb_extras(estac, bod2)
    assert out2["packs"][0]["precio_con_iva_uf"] is None
    # Hoja PACKS con la columna
    packs = [{"Pack Número": "P9", "Estacionamientos": "E-2A", "Bodegas": "B-1A",
              "PrecioUF": 240, "PrecioConIvaUF": 280}]
    assert _build_jb_extras(estac, bod, packs)["packs"][0]["precio_con_iva_uf"] == 280
    packs_sin = [{"Pack Número": "P9", "Estacionamientos": "E-2A", "Bodegas": "B-1A", "PrecioUF": 240}]
    assert "precio_con_iva_uf" not in _build_jb_extras(estac, bod, packs_sin)["packs"][0]


def test_columna_precio_en_revision_ya_no_existe_y_se_ignora(client, admin, db):
    """Nunca «precio en revisión» (29-oct): una columna así en el Excel no hace nada."""
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 2704.56, 3168, 1]],
                                  cols=COLS_SIN_IVA + ("PrecioEnRevision",)))
    u = _unidad(db, pid, "16B")
    assert u.precio_lista_uf == 2704.56 and u.precio_con_iva_uf == 3168
    assert not hasattr(Unidad, "precio_en_revision")
    assert "precio_en_revision" not in _unidad_dict(u)


def test_listado_desde_con_iva(client, admin, db):
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel(
        [["16B", "B1", 2704.56, 3168], ["17B", "B1", 2900, 3400]],
        cols=COLS_SIN_IVA))
    r = client.get("/proyectos", headers=h)
    fila = next(x for x in r.json() if x["id"] == pid)
    assert fila["precio_desde_uf"] == 2704.56
    assert fila["precio_desde_con_iva_uf"] == 3168


def test_listado_desde_con_iva_empate_no_depende_del_orden(client, admin, db):
    """Dos unidades con el mismo sin IVA: el «desde» normal es el de la que lo trae
    (y el menor), sin importar en qué orden devuelva las filas la base."""
    _, h = admin
    for orden in ([["16B", "B1", 2704.56, None], ["17B", "B1", 2704.56, 3170]],
                  [["17B", "B1", 2704.56, 3170], ["16B", "B1", 2704.56, None]],
                  [["16B", "B1", 2704.56, 3190], ["17B", "B1", 2704.56, 3170]]):
        pid = _proy(db)
        _subir(client, h, pid, _excel(orden, cols=COLS_SIN_IVA))
        fila = next(x for x in client.get("/proyectos", headers=h).json() if x["id"] == pid)
        assert fila["precio_desde_uf"] == 2704.56
        assert fila["precio_desde_con_iva_uf"] == 3170


# ── Encabezado ausente conserva; presente y vacío limpia ───────────────────────

def test_encabezado_ausente_conserva_y_vacio_limpia(client, admin, db):
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 2704.56, 3168]], cols=COLS_SIN_IVA))
    # Excel viejo / restauración por curl: sin la columna → no se borra nada
    _subir(client, h, pid, _excel([["16B", "B1", 2704.56]]))
    u = _unidad(db, pid, "16B")
    assert u.precio_con_iva_uf == 3168
    # Columna presente y vacía → se limpia
    _subir(client, h, pid, _excel([["16B", "B1", 3168, None]], cols=COLS_SIN_IVA))
    u = _unidad(db, pid, "16B")
    assert u.precio_con_iva_uf is None
    assert u.precio_lista_uf == 3168


# ── El editor no borra ─────────────────────────────────────────────────────────

def test_put_del_editor_de_unidad_no_borra_precio_normal(client, admin, db):
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 2704.56, 3168]], cols=COLS_SIN_IVA))
    u = _unidad(db, pid, "16B")
    copia_vieja = {"numero": "16B", "modelo": "B1", "tipo": "Depto", "precio_lista_uf": 2704.56,
                   "disponible": True, "precio_con_iva_uf": None}
    r = client.put(f"/proyectos/{pid}/unidades/{u.id}", headers=h, json=copia_vieja)
    assert r.status_code == 200, r.text
    u = _unidad(db, pid, "16B")
    assert u.precio_con_iva_uf == 3168
    assert r.json()["precio_con_iva_uf"] == 3168


def _put_proyecto(client, h, pid, extra):
    r = client.get(f"/proyectos/{pid}", headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    body["extra"] = extra
    r = client.put(f"/proyectos/{pid}", headers=h, json=body)
    assert r.status_code == 200, r.text
    return r.json()


def test_put_del_editor_de_proyecto_no_borra_listas_crudas_jb(client, admin, db):
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 2704.56, 3168]], cols=COLS_SIN_IVA,
                                  estac=[["E-1A", 173.11, 215]], bodegas=[["B-1A", 60, 75]],
                                  cols_sec=SEC_SIN_IVA))
    srv = _proyecto(db, pid).extra
    # Copia vieja del navegador: listas crudas sin el precio normal (o sin ellas)
    vieja = {k: v for k, v in srv.items() if k not in ("_bodegas_dom", "_excel_format")}
    vieja["_estacionamientos_dom"] = [{"Número": "E-1A", "PrecioUF": 215}]
    vieja["nota_editor"] = "editada a mano"
    _put_proyecto(client, h, pid, vieja)
    p = _proyecto(db, pid)
    assert p.extra["_estacionamientos_dom"] == srv["_estacionamientos_dom"]
    assert p.extra["_bodegas_dom"] == srv["_bodegas_dom"]
    assert p.extra["_excel_format"] == "jb_v2.4"
    assert p.extra["nota_editor"] == "editada a mano"  # el resto sí se guarda


def test_put_de_proyecto_sin_formato_jb_queda_como_siempre(client, admin, db):
    """Proyecto que nunca subió Excel JB: lo que manda el navegador se guarda tal cual."""
    _, h = admin
    pid = _proy(db, extra={"_estacionamientos_dom": [{"Número": "E-1"}]})
    _put_proyecto(client, h, pid, {"_estacionamientos_dom": [{"Número": "E-2"}]})
    assert _proyecto(db, pid).extra["_estacionamientos_dom"] == [{"Número": "E-2"}]


def test_put_borrado_explicito_en_proyecto_jb_sin_dato_queda_como_siempre(client, admin, db):
    """Proyecto JB sin precio normal: el PUT de scripts/wipe_proyecto_jb.py (extra sin las
    listas crudas) y el de scripts/patch_stock.py (pop de las listas) las borran, como
    antes del cambio."""
    _, h = admin
    # wipe modo full: extra = {jb_id}
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 3168]], estac=[["E-1A", 215]],
                                  bodegas=[["B-1A", 75]]))
    srv = _proyecto(db, pid).extra
    assert srv["_excel_format"] == "jb_v2.4" and srv["_estacionamientos_dom"]
    _put_proyecto(client, h, pid, {"jb_id": "abc"})
    ex = _proyecto(db, pid).extra
    for k in ("_estacionamientos_dom", "_bodegas_dom", "_excel_format"):
        assert k not in ex
    # wipe modo jb-only / patch_stock: borran las dos listas y dejan _excel_format
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 3168]], estac=[["E-1A", 215]],
                                  bodegas=[["B-1A", 75]]))
    srv = dict(_proyecto(db, pid).extra)
    srv.pop("_estacionamientos_dom")
    srv.pop("_bodegas_dom")
    srv["estacionamientos_dom"] = [{"numero": "E-9"}]
    _put_proyecto(client, h, pid, srv)
    ex = _proyecto(db, pid).extra
    assert "_estacionamientos_dom" not in ex and "_bodegas_dom" not in ex
    assert ex["_excel_format"] == "jb_v2.4"
    assert ex["estacionamientos_dom"] == [{"numero": "E-9"}]


# ── Timeline ───────────────────────────────────────────────────────────────────

def test_activacion_deja_un_solo_evento(client, admin, db):
    _, h = admin
    pid = _proy(db)
    filas_con = [["16B", "B1", 3168], ["17B", "B1", 3000], ["18B", "B1", 2900]]
    _subir(client, h, pid, _excel(filas_con, estac=[["E-1A", 215]]), origen="jb_importer")
    antes = list(_proyecto(db, pid).extra.get("timeline") or [])
    # Activación: ValorUF pasa a sin IVA y llega el precio normal
    _subir(client, h, pid, _excel(
        [["16B", "B1", 2704.56, 3168], ["17B", "B1", 2561, 3000],
         ["18B", "B1", 2476, 2900]],
        cols=COLS_SIN_IVA, estac=[["E-1A", 183.6, 215]], cols_sec=SEC_SIN_IVA), origen="jb_importer")
    tl = _proyecto(db, pid).extra["timeline"]
    nuevos = tl[:len(tl) - len(antes)]
    assert [e["tipo"] for e in nuevos] == ["Cambio a precio sin IVA"]
    assert "4 unidades: 3 deptos, 1 estacionamiento" in nuevos[0]["detalles"]
    assert _unidad(db, pid, "16B").precio_lista_uf == 2704.56
    # Corrida siguiente igual: sin cambios, sin eventos
    _subir(client, h, pid, _excel(
        [["16B", "B1", 2704.56, 3168], ["17B", "B1", 2561, 3000],
         ["18B", "B1", 2476, 2900]],
        cols=COLS_SIN_IVA, estac=[["E-1A", 183.6, 215]], cols_sec=SEC_SIN_IVA), origen="jb_importer")
    assert len(_proyecto(db, pid).extra["timeline"]) == len(tl)


def test_activado_cambio_de_estacionamiento_se_anota(client, admin, db):
    _, h = admin
    pid = _proy(db)
    filas = [["16B", "B1", 2704.56, 3168]]
    _subir(client, h, pid, _excel(filas, cols=COLS_SIN_IVA, estac=[["E-1A", 183.6, 215]],
                                  cols_sec=SEC_SIN_IVA), origen="jb_importer")
    n = len(_proyecto(db, pid).extra["timeline"])
    _subir(client, h, pid, _excel(filas, cols=COLS_SIN_IVA,
                                  estac=[["E-1A", 170, 200], ["E-2A", 150, 180]],
                                  cols_sec=SEC_SIN_IVA), origen="jb_importer")
    tl = _proyecto(db, pid).extra["timeline"]
    assert len(tl) == n + 1
    det = tl[0]["detalles"]
    assert tl[0]["tipo"] == "Excel Stock"
    assert "estacionamientos →" in det and "E-1A: precio 184 UF→170 UF" in det
    assert "precio normal 215 UF→200 UF" in det and "E-2A nuevo" in det


def test_cambio_de_precio_normal_de_un_depto_se_detalla(client, admin, db):
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 2704.56, 3168]], cols=COLS_SIN_IVA),
           origen="jb_importer")
    _subir(client, h, pid, _excel([["16B", "B1", 2650, 3100]], cols=COLS_SIN_IVA),
           origen="jb_importer")
    det = _proyecto(db, pid).extra["timeline"][0]["detalles"]
    assert "precio normal 3.168 UF→3.100 UF" in det
    assert "revisión" not in det


def test_reversion_deja_un_solo_evento(client, admin, db):
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 2704.56, 3168]], cols=COLS_SIN_IVA),
           origen="jb_importer")
    n = len(_proyecto(db, pid).extra["timeline"])
    # Interruptor apagado: el scraper escribe la columna vacía y ValorUF con IVA
    _subir(client, h, pid, _excel([["16B", "B1", 3168, None]], cols=COLS_SIN_IVA),
           origen="jb_importer")
    tl = _proyecto(db, pid).extra["timeline"]
    assert len(tl) == n + 1
    assert tl[0]["tipo"] == "Reversión a precio con IVA"
    assert "1 unidad: 1 depto" in tl[0]["detalles"]
    assert _unidad(db, pid, "16B").precio_lista_uf == 3168


# ── Plantilla de descarga ──────────────────────────────────────────────────────

def _hoja_unidad(client, h, pid):
    from openpyxl import load_workbook
    r = client.get(f"/proyectos/{pid}/unidades/excel/template?con_datos=true", headers=h)
    assert r.status_code == 200
    return list(load_workbook(io.BytesIO(r.content))["UNIDAD"].iter_rows(values_only=True))


def test_plantilla_sin_dato_tiene_las_columnas_de_siempre(client, admin, db):
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 3168]]))
    filas = _hoja_unidad(client, h, pid)
    assert "PrecioConIvaUF" not in filas[1] and len(filas[1]) == 17


def test_plantilla_con_dato_suma_columnas_y_vuelve_a_subir_igual(client, admin, db):
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 2704.56, 3168]], cols=COLS_SIN_IVA))
    filas = _hoja_unidad(client, h, pid)
    assert filas[1][-1] == "PrecioConIvaUF" and "PrecioEnRevision" not in filas[1]
    assert len(filas[1]) == 18
    assert filas[2][-1] == 3168


def test_activacion_parcial_cuenta_solo_las_que_traen_precio_normal(client, admin, db):
    """Activación donde un estacionamiento llega sin precio normal (celda vacía): el
    evento cuenta solo las unidades que de verdad pasaron a sin IVA."""
    _, h = admin
    pid = _proy(db)
    _subir(client, h, pid, _excel([["16B", "B1", 3168], ["17B", "B1", 3000]],
                                  estac=[["E-1A", 215], ["E-2A", 180]]), origen="jb_importer")
    antes = list(_proyecto(db, pid).extra.get("timeline") or [])
    _subir(client, h, pid, _excel(
        [["16B", "B1", 2704.56, 3168], ["17B", "B1", 2561, 3000]],
        cols=COLS_SIN_IVA, estac=[["E-1A", 183.6, 215], ["E-2A", 180, None]],
        cols_sec=SEC_SIN_IVA), origen="jb_importer")
    tl = _proyecto(db, pid).extra["timeline"]
    nuevos = tl[:len(tl) - len(antes)]
    act = next(e for e in nuevos if e["tipo"] == "Cambio a precio sin IVA")
    assert "3 unidades: 2 deptos, 1 estacionamiento" in act["detalles"]
    assert all("revisión" not in e["detalles"] for e in nuevos)
